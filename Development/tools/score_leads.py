#!/usr/bin/env python3
"""Auto-score leads and assign Tier based on Apeiron fit + data completeness.

Scoring criteria:
  - Apeiron sector fit (O&G, mining, chemicals, energy, etc.): +25
  - Apeiron role fit (instrumentation, automation, electrical, safety): +20
  - Has email: +15
  - Has LinkedIn: +10
  - Has phone: +5
  - Has Account_Summary (rich briefing): +5
  - Has Account_Size: +5
  - Has Account_City_State: +5
  - Priority = High: +10, Medium: +5

Tier assignment:
  - Tier 1 (score >= 70): High-value Apeiron-fit lead, ready for outreach
  - Tier 2 (score >= 45): Good potential, may need some enrichment
  - Tier 3 (score < 45): Low priority or incomplete

Usage:
    python score_leads.py              # preview scores
    python score_leads.py --save       # write Lead_Score + Tier to BD
"""
import argparse
import io
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

if sys.stdout and hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

APEIRON_SECTOR_RE = re.compile(
    r"oil|gas|o&g|petr[oó]leo|petroqui?m|energy|energia|mining|minera[cç][ãa]o|"
    r"siderur|chemicals?|qui?m|offshore|fpso|epc|refin|metalurgi|metals?|"
    r"celulose|papel|pulp|cimento|cement|port|terminal|utilit",
    re.IGNORECASE,
)
APEIRON_ROLE_RE = re.compile(
    r"instrument|automa[tcç]|el[eé]tric|electr|maintenance|manuten[cç]|"
    r"icss|scada|plc|clp|dcs|safety|seguran[cç]a|cybersec|ciberseg|"
    r"hmi|control|engenheiro|engineer",
    re.IGNORECASE,
)


def cell_text(ws, row, col):
    v = ws.cell(row, col).value
    if v is None:
        return ""
    return str(v).strip()


def score_lead(sector, role, email, linkedin, phone, summary, size, city, priority):
    """Calculate lead score (0-100)."""
    score = 0

    # Apeiron fit
    if sector and APEIRON_SECTOR_RE.search(sector):
        score += 25
    if role and APEIRON_ROLE_RE.search(role):
        score += 20

    # Contact channels
    if email:
        score += 15
    if linkedin:
        score += 10
    if phone:
        score += 5

    # Data quality
    if summary and len(summary) > 20:
        score += 5
    if size:
        score += 5
    if city:
        score += 5

    # Priority
    p = priority.lower() if priority else ""
    if "high" in p or "alta" in p:
        score += 10
    elif "medium" in p or "media" in p or "média" in p:
        score += 5

    return min(score, 100)


def assign_tier(score):
    if score >= 70:
        return "Tier 1"
    if score >= 45:
        return "Tier 2"
    return "Tier 3"


def main():
    parser = argparse.ArgumentParser(description="Score leads and assign tiers")
    parser.add_argument("--save", action="store_true", help="Write to BD")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[2]
    bd_path = root / "Apeiron_BR_Gestao_Comercial.xlsx"
    if not bd_path.exists():
        print(f"ERROR: BD not found: {bd_path}")
        return 1

    wb = openpyxl.load_workbook(bd_path)
    ws = wb["Leads"]

    headers = {}
    for col in range(1, ws.max_column + 1):
        h = cell_text(ws, 1, col)
        if h:
            headers[h] = col

    col = lambda name: headers.get(name)

    tiers = {"Tier 1": 0, "Tier 2": 0, "Tier 3": 0}
    scores = []
    sector_fit_count = 0
    role_fit_count = 0

    for row in range(2, ws.max_row + 1):
        sector = cell_text(ws, row, col("Account_Sector")) if col("Account_Sector") else ""
        role = cell_text(ws, row, col("Role")) if col("Role") else ""
        email = cell_text(ws, row, col("Email")) if col("Email") else ""
        linkedin = cell_text(ws, row, col("LinkedIn")) if col("LinkedIn") else ""
        phone = cell_text(ws, row, col("Account_Phone")) if col("Account_Phone") else ""
        summary = cell_text(ws, row, col("Account_Summary")) if col("Account_Summary") else ""
        size = cell_text(ws, row, col("Account_Size")) if col("Account_Size") else ""
        city = cell_text(ws, row, col("Account_City_State")) if col("Account_City_State") else ""
        priority = cell_text(ws, row, col("Priority")) if col("Priority") else ""

        s = score_lead(sector, role, email, linkedin, phone, summary, size, city, priority)
        t = assign_tier(s)
        tiers[t] += 1
        scores.append((row, s, t))

        if sector and APEIRON_SECTOR_RE.search(sector):
            sector_fit_count += 1
        if role and APEIRON_ROLE_RE.search(role):
            role_fit_count += 1

    total = len(scores)
    avg_score = sum(s for _, s, _ in scores) / total if total else 0

    print("=== Lead Scoring Report ===")
    print(f"Total leads: {total}")
    print(f"Average score: {avg_score:.1f}/100\n")

    for t in ["Tier 1", "Tier 2", "Tier 3"]:
        cnt = tiers[t]
        pct = cnt * 100 // total if total else 0
        print(f"  {t}: {cnt} ({pct}%)")

    print(f"\n--- Apeiron Fit ---")
    print(f"  Sector match: {sector_fit_count} ({sector_fit_count*100//total}%)")
    print(f"  Role match: {role_fit_count} ({role_fit_count*100//total}%)")

    # Score distribution
    print(f"\n--- Score Distribution ---")
    brackets = [(80, 100), (60, 79), (40, 59), (20, 39), (0, 19)]
    for lo, hi in brackets:
        cnt = sum(1 for _, s, _ in scores if lo <= s <= hi)
        print(f"  {lo}-{hi}: {cnt}")

    # Top 10
    top = sorted(scores, key=lambda x: x[1], reverse=True)
    print(f"\n--- Top 10 Leads ---")
    for row, s, t in top[:10]:
        name = cell_text(ws, row, col("Full Name")) if col("Full Name") else ""
        company = cell_text(ws, row, col("Company")) if col("Company") else ""
        print(f"  Score {s}: {name} ({company}) [{t}]")

    if args.save:
        backup_dir = root / "Support Files" / "Backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(bd_path, backup_dir / f"Apeiron_BR_Gestao_Comercial_backup_{ts}.xlsx")

        col_score = headers.get("Lead_Score")
        if not col_score:
            col_score = ws.max_column + 1
            ws.cell(1, col_score).value = "Lead_Score"

        col_tier = headers.get("Tier")
        if not col_tier:
            col_tier = ws.max_column + 1 if not col_score else ws.max_column + 1
            # Recalc in case Lead_Score was just added
            if col_tier == col_score:
                col_tier = col_score + 1
            ws.cell(1, col_tier).value = "Tier"

        for row, s, t in scores:
            ws.cell(row, col_score).value = s
            ws.cell(row, col_tier).value = t

        wb.save(bd_path)
        print(f"\nLead_Score + Tier saved to BD ({total} leads)")
    else:
        print("\n(use --save to write Lead_Score + Tier to BD)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
