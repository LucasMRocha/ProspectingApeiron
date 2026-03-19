#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook


def text(v) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    s = str(v).strip()
    if s.lower() in {"nan", "none", "nat", "<na>"}:
        return ""
    return s


def norm(v) -> str:
    return text(v).lower()


def slug(v) -> str:
    s = norm(v)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "unknown"


def clean_email(v) -> str:
    s = text(v).lower()
    if not s:
        return ""
    if "@" not in s or "." not in s.split("@")[-1]:
        return ""
    return s


def clean_phone(v) -> str:
    s = text(v)
    if not s:
        return ""
    digits = re.sub(r"[^\d+]", "", s)
    if len(re.sub(r"\D", "", digits)) < 8:
        return ""
    return digits


def clean_url(v) -> str:
    s = text(v)
    if not s:
        return ""
    if s.lower().startswith(("http://", "https://")):
        return s
    if "." in s:
        return f"https://{s}"
    return ""


def clean_cnpj(v) -> str:
    s = text(v)
    if not s:
        return ""
    digits = re.sub(r"\D", "", s)
    if len(digits) != 14:
        return ""
    return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"


def cnpj_digits(v) -> str:
    s = text(v)
    if not s:
        return ""
    digits = re.sub(r"\D", "", s)
    return digits if len(digits) == 14 else ""


def company_group_key(company: str, account_cnpj: str) -> str:
    digits = cnpj_digits(account_cnpj)
    if digits:
        return f"cnpj:{digits}"
    return f"name:{slug(company)}"


def to_int(v):
    s = text(v)
    if not s:
        return None
    try:
        return int(float(s))
    except Exception:
        return None


def find_project_root() -> Path:
    here = Path(__file__).resolve().parent
    for p in [here, *here.parents]:
        if (p / "Apeiron_BR_Gestao_Comercial.xlsx").exists():
            return p
    return Path(__file__).resolve().parents[1]


def backup_workbook(workbook_path: Path, backup_dir: Path, keep: int = 2) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    backup_path = backup_dir / f"Apeiron_BR_Gestao_Comercial_backup_{stamp}.xlsx"
    shutil.copyfile(workbook_path, backup_path)
    backups = sorted(
        backup_dir.glob("Apeiron_BR_Gestao_Comercial_backup_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    for old in backups[keep:]:
        try:
            old.unlink()
        except Exception:
            pass
    return backup_path


def header_map(ws, header_row: int) -> Dict[str, int]:
    out = {}
    for c in range(1, ws.max_column + 1):
        h = text(ws.cell(header_row, c).value)
        if h:
            out[h] = c
    return out


def rows_by_id(ws, headers: Dict[str, int], start_row: int) -> Dict[int, int]:
    out = {}
    id_col = headers.get("ID")
    if not id_col:
        return out
    for r in range(start_row, ws.max_row + 1):
        iid = to_int(ws.cell(r, id_col).value)
        if iid is None:
            continue
        out[iid] = r
    return out


def read_jsonl(path: Path) -> List[dict]:
    out: List[dict] = []
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            raw = line.strip()
            if not raw:
                continue
            try:
                obj = json.loads(raw)
                if isinstance(obj, dict):
                    out.append(obj)
            except Exception:
                print(f"Skipped invalid JSON at {path.name}:{line_no}")
    return out


def normalize_update_value(field: str, value: str) -> str:
    if field == "email":
        return clean_email(value)
    if field in {"phone", "account_phone"}:
        return clean_phone(value)
    if field in {"linkedin", "account_website"}:
        return clean_url(value)
    if field in {"account_cnpj", "cnpj"}:
        return clean_cnpj(value)
    return text(value)


def should_write(current_value: str, new_value: str, overwrite: bool) -> bool:
    if not text(new_value):
        return False
    if overwrite:
        return norm(current_value) != norm(new_value)
    return not text(current_value)


def update_cell(ws, row: int, col: int, new_value: str, overwrite: bool) -> bool:
    current = text(ws.cell(row, col).value)
    if not should_write(current, new_value, overwrite):
        return False
    ws.cell(row, col).value = new_value
    return True


def export_leads_js(workbook_path: Path, out_js: Path) -> int:
    xls = pd.ExcelFile(workbook_path)
    ops = pd.read_excel(xls, sheet_name="Leads_Operations")
    det = pd.read_excel(xls, sheet_name="Leads_Details")
    det_map = {int(r["ID"]): r for _, r in det.dropna(subset=["ID"]).iterrows()}
    rows = []
    for _, r in ops.dropna(subset=["ID"]).iterrows():
        iid = int(r["ID"])
        d = det_map.get(iid, {})
        item = {
            "id": iid,
            "company": text(r.get("Company")),
            "name": text(r.get("Full Name")),
            "role": text(r.get("Role")),
            "status": text(r.get("Status")) or text(d.get("Status")),
            "priority": text(r.get("Priority")) or text(d.get("Priority")),
            "nextAction": text(r.get("Next Action")) or text(d.get("Next Action")),
            "nextActionDate": text(r.get("Next Action Date")) or text(d.get("Next Action Date")),
            "channel": text(r.get("Channel")),
            "owner": text(r.get("Lead Insight")) or text(d.get("Lead Insight")),
            "source": text(r.get("Source")) or text(d.get("Source")),
            "sharePointId": text(r.get("SharePoint_ID")) or text(d.get("SharePoint_ID")),
            "possibleDuplicate": text(r.get("Possible Duplicate")),
            "email": text(d.get("Email")),
            "phone": text(d.get("Phone")),
            "linkedin": text(d.get("LinkedIn")),
            "location": text(d.get("Location")),
            "department": text(d.get("Department")),
            "notes": text(d.get("Notes")),
            "accountSummary": text(d.get("Account_Summary")),
            "accountSector": text(d.get("Account_Sector")),
            "accountCityState": text(d.get("Account_City_State")),
            "accountSize": text(d.get("Account_Size")),
            "accountRevenue": text(d.get("Account_Revenue")),
            "accountWebsite": text(d.get("Account_Website")),
            "accountPhone": text(d.get("Account_Phone")),
            "accountCnpj": text(d.get("Account_CNPJ")),
            "accountBrand": text(d.get("Account_Brand") or d.get("Brand")),
        }
        rows.append(item)
    out_js.parent.mkdir(parents=True, exist_ok=True)
    out_js.write_text("window.INITIAL_DATA = " + json.dumps(rows, ensure_ascii=False) + ";\n", encoding="utf-8")
    return len(rows)


def build_company_id_map(wb) -> Tuple[Dict[str, str], Dict[str, List[int]]]:
    ws_ops = wb["Leads_Operations"]
    ws_det = wb["Leads_Details"]
    h_ops = header_map(ws_ops, 1)
    h_det = header_map(ws_det, 1)

    id_col_ops = h_ops.get("ID")
    id_col_det = h_det.get("ID")
    if not id_col_ops or not id_col_det:
        return {}, {}

    id_to_company: Dict[int, str] = {}
    company_col_ops = h_ops.get("Company")
    if company_col_ops:
        for r in range(2, ws_ops.max_row + 1):
            iid = to_int(ws_ops.cell(r, id_col_ops).value)
            if iid is None:
                continue
            id_to_company[iid] = text(ws_ops.cell(r, company_col_ops).value)

    grouped: Dict[str, List[int]] = defaultdict(list)
    cnpj_col_det = h_det.get("Account_CNPJ")
    company_col_det = h_det.get("Company")
    for r in range(2, ws_det.max_row + 1):
        iid = to_int(ws_det.cell(r, id_col_det).value)
        if iid is None:
            continue
        company = text(ws_det.cell(r, company_col_det).value) if company_col_det else ""
        if not company:
            company = id_to_company.get(iid, "")
        cnpj = text(ws_det.cell(r, cnpj_col_det).value) if cnpj_col_det else ""
        key = company_group_key(company, cnpj)
        grouped[key].append(iid)

    keys = sorted(grouped.keys())
    key_to_company_id = {k: f"C{idx + 1:04d}" for idx, k in enumerate(keys)}
    company_id_to_ids = {key_to_company_id[k]: sorted(set(v)) for k, v in grouped.items()}
    return key_to_company_id, company_id_to_ids


def build_company_name_to_ids_map(wb) -> Dict[str, List[int]]:
    ws_det = wb["Leads_Details"]
    h_det = header_map(ws_det, 1)
    c_id = h_det.get("ID")
    c_company = h_det.get("Company")
    if not c_id or not c_company:
        return {}
    out: Dict[str, List[int]] = defaultdict(list)
    for r in range(2, ws_det.max_row + 1):
        iid = to_int(ws_det.cell(r, c_id).value)
        company = text(ws_det.cell(r, c_company).value)
        if iid is None or not company:
            continue
        out[norm(company)].append(iid)
    return {k: sorted(set(v)) for k, v in out.items()}


def apply_company_results(
    wb,
    company_results: List[dict],
    company_id_to_ids: Dict[str, List[int]],
    company_group_to_ids: Dict[str, List[int]],
    company_name_to_ids: Dict[str, List[int]],
    id_to_row_det: Dict[int, int],
    headers_det: Dict[str, int],
    overwrite: bool,
) -> Dict[str, int]:
    stats = defaultdict(int)
    field_map = {
        "account_sector": "Account_Sector",
        "account_city_state": "Account_City_State",
        "account_website": "Account_Website",
        "account_phone": "Account_Phone",
        "account_size": "Account_Size",
        "account_revenue": "Account_Revenue",
        "account_summary": "Account_Summary",
        "account_cnpj": "Account_CNPJ",
        "account_brand": "Account_Brand",
    }
    ws_det = wb["Leads_Details"]

    for item in company_results:
        company_id = text(item.get("company_id"))
        company = text(item.get("company"))
        group_key = text(item.get("group_key"))

        ids: List[int] = []
        if company_id:
            ids = company_id_to_ids.get(company_id, [])
        elif group_key:
            ids = company_group_to_ids.get(group_key, [])
        elif company:
            ids = company_name_to_ids.get(norm(company), [])
        if not ids:
            continue
        updates = item.get("updates", {})
        if not isinstance(updates, dict):
            continue
        for iid in ids:
            row = id_to_row_det.get(iid)
            if not row:
                continue
            for key, val in updates.items():
                target_col_name = field_map.get(key)
                if not target_col_name:
                    continue
                col = headers_det.get(target_col_name)
                if not col:
                    continue
                normalized = normalize_update_value(key, val)
                if update_cell(ws_det, row, col, normalized, overwrite):
                    stats[f"company:{key}"] += 1
                    stats["company:total_cells"] += 1
    return dict(stats)


def apply_contact_results(
    wb,
    contact_results: List[dict],
    id_to_row_ops: Dict[int, int],
    id_to_row_det: Dict[int, int],
    headers_ops: Dict[str, int],
    headers_det: Dict[str, int],
    overwrite: bool,
) -> Dict[str, int]:
    stats = defaultdict(int)
    field_map_det = {
        "role": "Role",
        "department": "Department",
        "email": "Email",
        "phone": "Phone",
        "linkedin": "LinkedIn",
        "location": "Location",
    }
    field_map_ops = {
        "role": "Role",
    }
    ws_ops = wb["Leads_Operations"]
    ws_det = wb["Leads_Details"]

    for item in contact_results:
        iid = to_int(item.get("id"))
        if iid is None:
            continue
        updates = item.get("updates", {})
        if not isinstance(updates, dict):
            continue
        row_ops = id_to_row_ops.get(iid)
        row_det = id_to_row_det.get(iid)
        if not row_ops and not row_det:
            continue

        for key, val in updates.items():
            normalized = normalize_update_value(key, val)
            det_col_name = field_map_det.get(key)
            if row_det and det_col_name and headers_det.get(det_col_name):
                if update_cell(ws_det, row_det, headers_det[det_col_name], normalized, overwrite):
                    stats[f"contact:{key}"] += 1
                    stats["contact:total_cells"] += 1
            ops_col_name = field_map_ops.get(key)
            if row_ops and ops_col_name and headers_ops.get(ops_col_name):
                if update_cell(ws_ops, row_ops, headers_ops[ops_col_name], normalized, overwrite):
                    stats[f"contact:{key}:ops"] += 1
    return dict(stats)


def sync_accounts_from_details(wb) -> Dict[str, int]:
    if "Accounts" not in wb.sheetnames or "Leads_Details" not in wb.sheetnames:
        return {}

    ws_acc = wb["Accounts"]
    ws_det = wb["Leads_Details"]
    h_acc = header_map(ws_acc, 1)
    h_det = header_map(ws_det, 1)
    if not h_acc or not h_det:
        return {}

    c_company_acc = h_acc.get("Company")
    if not c_company_acc:
        return {}

    c_company_det = h_det.get("Company")
    if not c_company_det:
        return {}

    det_by_company: Dict[str, Dict[str, str]] = {}
    for r in range(2, ws_det.max_row + 1):
        company = text(ws_det.cell(r, c_company_det).value)
        if not company:
            continue
        key = norm(company)
        cur = det_by_company.setdefault(
            key,
            {
                "Sector": "",
                "City/State": "",
                "Website": "",
                "Phone": "",
                "CNPJ": "",
                "Brand": "",
            },
        )
        vals = {
            "Sector": text(ws_det.cell(r, h_det.get("Account_Sector")).value) if h_det.get("Account_Sector") else "",
            "City/State": text(ws_det.cell(r, h_det.get("Account_City_State")).value) if h_det.get("Account_City_State") else "",
            "Website": text(ws_det.cell(r, h_det.get("Account_Website")).value) if h_det.get("Account_Website") else "",
            "Phone": text(ws_det.cell(r, h_det.get("Account_Phone")).value) if h_det.get("Account_Phone") else "",
            "CNPJ": text(ws_det.cell(r, h_det.get("Account_CNPJ")).value) if h_det.get("Account_CNPJ") else "",
            "Brand": text(ws_det.cell(r, h_det.get("Account_Brand")).value) if h_det.get("Account_Brand") else "",
        }
        for k, v in vals.items():
            if not cur[k] and text(v):
                cur[k] = text(v)

    stats = defaultdict(int)
    for r in range(2, ws_acc.max_row + 1):
        company = text(ws_acc.cell(r, c_company_acc).value)
        if not company:
            continue
        payload = det_by_company.get(norm(company))
        if not payload:
            continue
        for acc_col, val in payload.items():
            c = h_acc.get(acc_col)
            if not c or not text(val):
                continue
            if update_cell(ws_acc, r, c, val, overwrite=True):
                stats[f"accounts:{acc_col}"] += 1
                stats["accounts:total_cells"] += 1
    return dict(stats)


def maybe_build_single_file(root: Path, do_build: bool) -> None:
    if not do_build:
        return
    build_script = root / "Development" / "tools" / "build-single-file.ps1"
    if not build_script.exists():
        return
    try:
        subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(build_script),
            ],
            check=True,
            cwd=str(root),
        )
    except Exception as exc:
        print(f"Warning: build-single-file failed: {exc}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Merge enrichment JSONL results back into DB.")
    parser.add_argument("--workbook", type=Path, default=None, help="Path to Apeiron_BR_Gestao_Comercial.xlsx")
    parser.add_argument("--company-results", type=Path, default=None, help="Path to company_enrichment_results.jsonl")
    parser.add_argument("--contact-results", type=Path, default=None, help="Path to contact_enrichment_results.jsonl")
    parser.add_argument("--overwrite", action="store_true", help="Allow overwriting non-empty cells")
    parser.add_argument("--dry-run", action="store_true", help="Calculate updates without writing workbook")
    parser.add_argument("--skip-build", action="store_true", help="Skip rebuilding single-file dashboard")
    args = parser.parse_args()

    root = find_project_root()
    support = root / "Support Files"
    workbook = args.workbook or (root / "Apeiron_BR_Gestao_Comercial.xlsx")
    enrichment_dir = support / "_output" / "enrichment"
    company_results_path = args.company_results or (enrichment_dir / "company_enrichment_results.jsonl")
    contact_results_path = args.contact_results or (enrichment_dir / "contact_enrichment_results.jsonl")

    company_results = read_jsonl(company_results_path)
    contact_results = read_jsonl(contact_results_path)
    if not company_results and not contact_results:
        raise SystemExit("No enrichment results found. Nothing to merge.")

    backup = None
    if not args.dry_run:
        backup = backup_workbook(workbook, support / "_backup", keep=2)
    wb = load_workbook(workbook)
    ws_ops = wb["Leads_Operations"]
    ws_det = wb["Leads_Details"]
    headers_ops = header_map(ws_ops, 1)
    headers_det = header_map(ws_det, 1)
    id_to_row_ops = rows_by_id(ws_ops, headers_ops, 2)
    id_to_row_det = rows_by_id(ws_det, headers_det, 2)

    key_to_company_id, company_id_to_ids = build_company_id_map(wb)
    company_group_to_ids = {k: company_id_to_ids.get(v, []) for k, v in key_to_company_id.items()}
    company_name_to_ids = build_company_name_to_ids_map(wb)

    stats = {}
    stats.update(
        apply_company_results(
            wb,
            company_results,
            company_id_to_ids,
            company_group_to_ids,
            company_name_to_ids,
            id_to_row_det,
            headers_det,
            args.overwrite,
        )
    )
    stats.update(
        apply_contact_results(
            wb,
            contact_results,
            id_to_row_ops,
            id_to_row_det,
            headers_ops,
            headers_det,
            args.overwrite,
        )
    )
    stats.update(sync_accounts_from_details(wb))

    rows_exported = 0
    if not args.dry_run:
        wb.save(workbook)
        rows_exported = export_leads_js(workbook, root / "Development" / "src" / "dashboard" / "data" / "leads.js")
        maybe_build_single_file(root, do_build=not args.skip_build)

    report = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(workbook),
        "backup_created": str(backup) if backup else "",
        "company_results_file": str(company_results_path),
        "contact_results_file": str(contact_results_path),
        "company_results_count": len(company_results),
        "contact_results_count": len(contact_results),
        "rows_exported_to_leads_js": rows_exported,
        "overwrite_mode": bool(args.overwrite),
        "dry_run": bool(args.dry_run),
        "stats": stats,
    }
    out_dir = support / "_output" / "enrichment"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_json = out_dir / "merge_enrichment_report.json"
    out_md = out_dir / "merge_enrichment_report.md"
    out_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    md_lines = [
        "# Enrichment Merge Report",
        "",
        f"- Timestamp: {report['timestamp']}",
        f"- Workbook: `{report['workbook']}`",
        f"- Dry run: `{report['dry_run']}`",
        f"- Backup created: `{report['backup_created'] or 'not created (dry-run)'}`",
        f"- Company result lines: {report['company_results_count']}",
        f"- Contact result lines: {report['contact_results_count']}",
        f"- Rows exported to `leads.js`: {report['rows_exported_to_leads_js']}",
        f"- Overwrite mode: `{report['overwrite_mode']}`",
        "",
        "## Cell Updates",
    ]
    if stats:
        for k in sorted(stats.keys()):
            md_lines.append(f"- {k}: {stats[k]}")
    else:
        md_lines.append("- No cell updates were applied.")
    out_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
