
#!/usr/bin/env python3
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import urllib.error
import urllib.parse
import urllib.request
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl


def text(v: Any) -> str:
    return "" if v is None else str(v)


def norm(v: Any) -> str:
    return text(v).strip().lower()


def to_int(v: Any):
    s = text(v).strip()
    if not s:
        return None
    try:
        return int(float(s))
    except Exception:
        return None


def to_float(v: Any):
    s = text(v).strip().replace(",", ".")
    if not s:
        return None
    try:
        x = float(s)
        if x != x:
            return None
        return x
    except Exception:
        return None


def first_non_empty(*vals: Any) -> str:
    for v in vals:
        s = text(v).strip()
        if s:
            return s
    return ""


def safe_name(v: Any) -> str:
    if isinstance(v, dict):
        return text(v.get("name") or v.get("display_value") or "").strip()
    return text(v).strip()


def record_get(rec: dict[str, Any], *keys: str):
    for k in keys:
        if k in rec:
            return rec[k]
    low = {k.lower(): k for k in rec.keys()}
    for k in keys:
        kk = low.get(k.lower())
        if kk is not None:
            return rec[kk]
    return None


def parse_date(v: Any):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = text(v).strip()
    if not s:
        return None
    s = s.replace("Z", "+00:00")
    try:
        return dt.datetime.fromisoformat(s).date()
    except Exception:
        pass
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        try:
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except Exception:
            return None
    m = re.match(r"^(\d{2})/(\d{2})/(\d{4})$", s)
    if m:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except Exception:
            return None
    return None


def week_of(d):
    return d.isocalendar().week if d else None


def quarter_of(d):
    if not d:
        return ""
    q = ((d.month - 1) // 3) + 1
    return f"{str(d.year)[-2:]}_Q{q}"


def map_stage(raw: str) -> str:
    s = norm(raw)
    if not s:
        return "Lead"
    if any(x in s for x in ["won", "lost", "closed", "deal"]):
        return "Deal"
    if "negot" in s:
        return "Negotiation"
    if any(x in s for x in ["propos", "quote", "cot"]):
        return "Proposal"
    if any(x in s for x in ["qualif", "contact", "work"]):
        return "Qualification"
    return "Lead"


def map_status(raw: str) -> str:
    s = norm(raw)
    if not s:
        return "Open"
    if any(x in s for x in ["won", "ganho", "closed won"]):
        return "Closed - Won"
    if any(x in s for x in ["lost", "perd", "junk", "disqual", "closed lost"]):
        return "Closed - Lost"
    if "internal denial" in s:
        return "Internal denial"
    return "Open"


def map_urgency(raw: str) -> str:
    s = norm(raw)
    if any(x in s for x in ["high", "hot", "urgent", "alta"]):
        return "High"
    if any(x in s for x in ["low", "cold", "baixa"]):
        return "Low"
    if "no urgency" in s:
        return "No urgency"
    return "Medium"


def size_from_employees(v: Any, size_raw: Any):
    s = text(size_raw).strip()
    if s:
        return s
    n = to_int(v)
    if n is None:
        return ""
    if n >= 1000:
        return "Big"
    if n >= 200:
        return "Medium"
    return "Small"


def parse_prob(v: Any):
    p = to_float(v)
    if p is None:
        return None
    if p > 1:
        p = p / 100.0
    return max(0.0, min(1.0, p))


def by_key_contains(rec: dict[str, Any], *parts: str) -> str:
    parts = [x.lower() for x in parts]
    for k, v in rec.items():
        lk = k.lower()
        if any(p in lk for p in parts):
            s = text(v).strip()
            if s:
                return s
    return ""


def normalize_record(module: str, rec: dict[str, Any]) -> dict[str, Any]:
    zoho_id = text(rec.get("id")).strip()
    if not zoho_id:
        return {}

    company = first_non_empty(
        record_get(rec, "Company"),
        safe_name(record_get(rec, "Account_Name", "Account Name", "Account")),
        "Unknown Company",
    )
    first = text(record_get(rec, "First_Name", "First Name")).strip()
    last = text(record_get(rec, "Last_Name", "Last Name")).strip()
    contact = first_non_empty(
        record_get(rec, "Full_Name", "Full Name"),
        safe_name(record_get(rec, "Contact_Name", "Contact Name")),
        " ".join([x for x in [first, last] if x]).strip(),
        safe_name(record_get(rec, "Name")),
        "Unnamed Contact",
    )

    created = parse_date(first_non_empty(record_get(rec, "Created_Time"), record_get(rec, "Created_Date"), record_get(rec, "Created")))
    stage_date = parse_date(first_non_empty(record_get(rec, "Modified_Time"), record_get(rec, "Updated_Time")))
    last_contact = parse_date(first_non_empty(record_get(rec, "Last_Activity_Time"), record_get(rec, "Last_Contacted_Time")))
    next_action = parse_date(first_non_empty(record_get(rec, "Next_Activity_Date"), record_get(rec, "Follow_Up_Date"), record_get(rec, "Closing_Date"), record_get(rec, "Expected_Closing_Date")))

    stage_raw = first_non_empty(record_get(rec, "Stage"), record_get(rec, "Lead_Status"), record_get(rec, "Pipeline_Stage"))
    status_raw = first_non_empty(record_get(rec, "Status"), record_get(rec, "Lead_Status"), stage_raw)
    status = map_status(status_raw)

    prob = parse_prob(first_non_empty(record_get(rec, "Probability"), record_get(rec, "Deal_Probability")))
    est = to_float(first_non_empty(record_get(rec, "Estimated_Value"), record_get(rec, "Amount"), record_get(rec, "Expected_Revenue")))
    fval = to_float(first_non_empty(record_get(rec, "Forecast_Deal_Value"), record_get(rec, "Forecast_Value")))
    if fval is None and est is not None and prob is not None:
        fval = round(est * prob, 2)

    return {
        "zoho_key": f"{module}:{zoho_id}",
        "zoho_module": module,
        "Description": first_non_empty(record_get(rec, "Description"), record_get(rec, "Deal_Name"), record_get(rec, "Subject")),
        "Date_Creation": created,
        "Week": week_of(created),
        "Source_Lead": first_non_empty(record_get(rec, "Lead_Source"), record_get(rec, "Source"), f"Zoho CRM ({module})"),
        "Company": company,
        "Size_Company": size_from_employees(first_non_empty(record_get(rec, "No_of_Employees"), record_get(rec, "Employees")), record_get(rec, "Size_Company")),
        "Contact_Company": contact,
        "Tax_ID": first_non_empty(record_get(rec, "Tax_ID", "Tax ID"), by_key_contains(rec, "cnpj"), by_key_contains(rec, "tax id"), by_key_contains(rec, "tax_id")),
        "Segment": first_non_empty(record_get(rec, "Industry"), record_get(rec, "Segment")),
        "Country": first_non_empty(record_get(rec, "Country"), "Brasil"),
        "Estate": first_non_empty(record_get(rec, "State")),
        "City": first_non_empty(record_get(rec, "City")),
        "Seller": safe_name(record_get(rec, "Owner", "Record Owner", "Sales_Owner")),
        "Office": first_non_empty(record_get(rec, "Office"), record_get(rec, "Territory"), record_get(rec, "Region"), "Brazil"),
        "Funnel_Stage": map_stage(stage_raw),
        "Status": status,
        "Reason_Lost": first_non_empty(record_get(rec, "Reason_for_Loss"), record_get(rec, "Loss_Reason"), "NA") if status in ("Closed - Lost", "Internal denial") else "NA",
        "Date_Actual_Stage": stage_date,
        "Days_on_Stage": (dt.date.today() - stage_date).days if stage_date else None,
        "Date_Last_Contact": last_contact,
        "Days_Since_Last_Contact": (dt.date.today() - last_contact).days if last_contact else None,
        "Date_Next_Action": next_action,
        "OBS": first_non_empty(record_get(rec, "Description"), record_get(rec, "Next_Step"), record_get(rec, "Next_Action"), f"Imported from Zoho CRM ({module})"),
        "RelationShip_With_Customer": first_non_empty(record_get(rec, "RelationShip_With_Customer"), record_get(rec, "Relationship_With_Customer"), "New Client"),
        "Urgency": map_urgency(first_non_empty(record_get(rec, "Rating"), record_get(rec, "Priority"), record_get(rec, "Urgency"))),
        "Technical_Fit": first_non_empty(record_get(rec, "Technical_Fit"), record_get(rec, "Fit")),
        "Budget": first_non_empty(record_get(rec, "Budget"), record_get(rec, "Budget_Status")),
        "Probability": prob,
        "Type of service": first_non_empty(record_get(rec, "Type of service"), record_get(rec, "Type_of_Service"), record_get(rec, "Service_Type")),
        "Currency": first_non_empty(record_get(rec, "Currency"), "BRL"),
        "Forecast _Date": first_non_empty(record_get(rec, "Forecast _Date"), record_get(rec, "Forecast_Date"), quarter_of(next_action or created)),
        "Estimated_Value": est,
        "Forecast_Deal_Value": fval,
        "Contact_Email": first_non_empty(record_get(rec, "Email"), by_key_contains(rec, "email")),
        "Contact_Phone": first_non_empty(record_get(rec, "Phone"), record_get(rec, "Mobile"), by_key_contains(rec, "phone"), by_key_contains(rec, "mobile")),
        "Contact_LinkedIn": first_non_empty(record_get(rec, "LinkedIn", "Linkedin"), by_key_contains(rec, "linkedin")),
        "Contact_WhatsApp": first_non_empty(record_get(rec, "WhatsApp"), by_key_contains(rec, "whatsapp", "whats app", "what's app")),
    }

def http_json(method: str, url: str, headers: dict[str, str] | None = None, form: dict[str, str] | None = None) -> dict[str, Any]:
    hdr = headers or {}
    data = None
    if form is not None:
        data = urllib.parse.urlencode(form).encode("utf-8")
        hdr = {**hdr, "Content-Type": "application/x-www-form-urlencoded"}
    req = urllib.request.Request(url, headers=hdr, data=data, method=method.upper())
    try:
        with urllib.request.urlopen(req, timeout=45) as resp:
            body = resp.read().decode("utf-8", errors="replace")
            return json.loads(body) if body else {}
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {e.code} calling {url}: {body}") from e
    except urllib.error.URLError as e:
        raise RuntimeError(f"Network error calling {url}: {e}") from e


class ZohoClient:
    def __init__(self, cid: str, secret: str, refresh: str, accounts_domain: str, api_domain: str = ""):
        self.cid = cid
        self.secret = secret
        self.refresh = refresh
        self.accounts_domain = accounts_domain
        self.api_domain_override = api_domain
        self.token = ""
        self.api_domain = ""

    def authenticate(self):
        data = http_json(
            "POST",
            f"https://{self.accounts_domain}/oauth/v2/token",
            form={
                "grant_type": "refresh_token",
                "client_id": self.cid,
                "client_secret": self.secret,
                "refresh_token": self.refresh,
            },
        )
        token = text(data.get("access_token")).strip()
        if not token:
            raise RuntimeError(f"Zoho auth failed: {data}")
        self.token = token
        self.api_domain = self.api_domain_override.strip() or text(data.get("api_domain")).strip() or "https://www.zohoapis.com"

    def get_records(self, module: str, max_pages: int, per_page: int) -> list[dict[str, Any]]:
        if not self.token:
            self.authenticate()
        out: list[dict[str, Any]] = []
        page = 1
        while page <= max_pages:
            url = f"{self.api_domain.rstrip('/')}/crm/v8/{urllib.parse.quote(module)}?page={page}&per_page={per_page}"
            data = http_json("GET", url, headers={"Authorization": f"Zoho-oauthtoken {self.token}"})
            rows = data.get("data") or []
            if not isinstance(rows, list):
                break
            out.extend(rows)
            info = data.get("info") or {}
            if not info.get("more_records") or not rows:
                break
            page += 1
        return out


def ensure_col(ws, header_row: int, name: str) -> int:
    for c in range(1, ws.max_column + 1):
        if text(ws.cell(header_row, c).value).strip() == name:
            return c
    c = ws.max_column + 1
    ws.cell(header_row, c).value = name
    return c


def load_cols(ws, header_row: int) -> dict[str, int]:
    cols: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = text(ws.cell(header_row, c).value).strip()
        if h:
            cols[h] = c
    return cols


def setv(ws, cols: dict[str, int], row: int, key: str, value: Any, overwrite: bool = False, force: bool = False):
    c = cols.get(key)
    if not c:
        return
    if value is None:
        return
    s = text(value).strip()
    if not force and s == "":
        return
    cur = text(ws.cell(row, c).value).strip()
    if force or overwrite or cur == "":
        ws.cell(row, c).value = value


def copy_style_row(ws, src: int, dst: int):
    if src in ws.row_dimensions and ws.row_dimensions[src].height is not None:
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height
    for c in range(1, ws.max_column + 1):
        a = ws.cell(src, c)
        b = ws.cell(dst, c)
        if a.has_style:
            b._style = copy(a._style)
        b.number_format = a.number_format
        b.value = None


def build_indexes(ws, cols: dict[str, int], data_start: int):
    by_id = {}
    by_zoho = {}
    by_tax_company = {}
    by_company_contact_desc = {}
    by_company_contact = {}
    max_id = 0
    last_row = data_start - 1

    idc = cols["ID_Opportunity"]
    for r in range(data_start, ws.max_row + 1):
        iid = to_int(ws.cell(r, idc).value)
        if iid is None:
            continue
        by_id[iid] = r
        max_id = max(max_id, iid)
        last_row = max(last_row, r)

        z = norm(ws.cell(r, cols.get("Zoho_Key", 0)).value if cols.get("Zoho_Key") else "")
        if z:
            by_zoho[z] = r

        tax = norm(ws.cell(r, cols.get("Tax_ID", 0)).value if cols.get("Tax_ID") else "")
        comp = norm(ws.cell(r, cols.get("Company", 0)).value if cols.get("Company") else "")
        con = norm(ws.cell(r, cols.get("Contact_Company", 0)).value if cols.get("Contact_Company") else "")
        des = norm(ws.cell(r, cols.get("Description", 0)).value if cols.get("Description") else "")
        if tax and comp:
            by_tax_company[(tax, comp)] = r
        if comp and con and des:
            by_company_contact_desc[(comp, con, des)] = r
        if comp and con:
            by_company_contact[(comp, con)] = r

    return {
        "by_id": by_id,
        "by_zoho": by_zoho,
        "by_tax_company": by_tax_company,
        "by_company_contact_desc": by_company_contact_desc,
        "by_company_contact": by_company_contact,
        "max_id": max_id,
        "last_row": last_row,
    }

def find_match_row(rec: dict[str, Any], idx: dict[str, Any]):
    z = norm(rec.get("zoho_key"))
    if z and z in idx["by_zoho"]:
        return idx["by_zoho"][z]
    tax = norm(rec.get("Tax_ID"))
    comp = norm(rec.get("Company"))
    con = norm(rec.get("Contact_Company"))
    des = norm(rec.get("Description"))
    if tax and comp and (tax, comp) in idx["by_tax_company"]:
        return idx["by_tax_company"][(tax, comp)]
    if comp and con and des and (comp, con, des) in idx["by_company_contact_desc"]:
        return idx["by_company_contact_desc"][(comp, con, des)]
    if comp and con and (comp, con) in idx["by_company_contact"]:
        return idx["by_company_contact"][(comp, con)]
    return None


def refresh_indexes(row: int, ws, cols: dict[str, int], idx: dict[str, Any]):
    iid = to_int(ws.cell(row, cols["ID_Opportunity"]).value)
    if iid is not None:
        idx["by_id"][iid] = row
        idx["max_id"] = max(idx["max_id"], iid)
    z = norm(ws.cell(row, cols.get("Zoho_Key", 0)).value if cols.get("Zoho_Key") else "")
    if z:
        idx["by_zoho"][z] = row
    tax = norm(ws.cell(row, cols.get("Tax_ID", 0)).value if cols.get("Tax_ID") else "")
    comp = norm(ws.cell(row, cols.get("Company", 0)).value if cols.get("Company") else "")
    con = norm(ws.cell(row, cols.get("Contact_Company", 0)).value if cols.get("Contact_Company") else "")
    des = norm(ws.cell(row, cols.get("Description", 0)).value if cols.get("Description") else "")
    if tax and comp:
        idx["by_tax_company"][(tax, comp)] = row
    if comp and con and des:
        idx["by_company_contact_desc"][(comp, con, des)] = row
    if comp and con:
        idx["by_company_contact"][(comp, con)] = row


def sync_records(ws, cols: dict[str, int], data_start: int, records: list[dict[str, Any]], overwrite: bool):
    idx = build_indexes(ws, cols, data_start)
    created = 0
    updated = 0
    skipped = 0

    fixed = [
        "Description", "Date_Creation", "Week", "Source_Lead", "Company", "Size_Company", "Contact_Company", "Tax_ID",
        "Segment", "Country", "Estate", "City", "Seller", "Office", "Funnel_Stage", "Status", "Reason_Lost",
        "Date_Actual_Stage", "Days_on_Stage", "Date_Last_Contact", "Days_Since_Last_Contact", "Date_Next_Action", "OBS",
        "RelationShip_With_Customer", "Urgency", "Technical_Fit", "Budget", "Probability", "Type of service", "Currency",
        "Forecast _Date", "Estimated_Value", "Forecast_Deal_Value", "Contact_Email", "Contact_Phone", "Contact_LinkedIn", "Contact_WhatsApp"
    ]

    for rec in records:
        if not rec:
            skipped += 1
            continue

        row = find_match_row(rec, idx)
        is_new = row is None
        if is_new:
            row = max(idx["last_row"] + 1, data_start)
            src = row - 1 if row > data_start else row
            if src >= data_start:
                copy_style_row(ws, src, row)
            idx["max_id"] += 1
            ws.cell(row, cols["ID_Opportunity"]).value = idx["max_id"]
            idx["last_row"] = row
            created += 1
        else:
            updated += 1

        ow = overwrite or is_new
        for k in fixed:
            setv(ws, cols, row, k, rec.get(k), overwrite=ow)

        setv(ws, cols, row, "Zoho_Key", rec.get("zoho_key"), force=True)
        setv(ws, cols, row, "Zoho_Module", rec.get("zoho_module"), force=True)
        setv(ws, cols, row, "Zoho_Last_Sync", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), force=True)
        refresh_indexes(row, ws, cols, idx)

    return {"created": created, "updated": updated, "skipped": skipped, "total": len(records)}


def serialize(v: Any):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and (v != v or v == float("inf") or v == float("-inf")):
        return None
    if isinstance(v, str):
        d = parse_date(v)
        if d:
            return d.isoformat()
    return v


def export_leads_js(workbook: Path, sheet: str, header_row: int, data_start: int, out_js: Path):
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb[sheet]
    cols = load_cols(ws, header_row)
    idc = cols["ID_Opportunity"]
    rows = []
    for r in range(data_start, ws.max_row + 1):
        iid = to_int(ws.cell(r, idc).value)
        if iid is None:
            continue
        item = {}
        for h, c in cols.items():
            item[h] = serialize(ws.cell(r, c).value)
        rows.append(item)
    out_js.parent.mkdir(parents=True, exist_ok=True)
    out_js.write_text("window.RAW_DATA = " + json.dumps(rows, ensure_ascii=False) + ";\n", encoding="utf-8")
    return len(rows)


def create_backup(workbook: Path, keep: int):
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = workbook.parent / f"Dashboard_Comercial_Brasil_backup_{ts}{workbook.suffix}"
    shutil.copy2(workbook, backup)
    all_bak = sorted(workbook.parent.glob("Dashboard_Comercial_Brasil_backup_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in all_bak[max(1, keep):]:
        try:
            old.unlink()
            print(f"Deleted old backup: {old.name}")
        except Exception as e:
            print(f"Warning deleting backup {old.name}: {e}")
    return backup


def need_env(name: str) -> str:
    val = os.getenv(name, "").strip()
    if not val:
        raise RuntimeError(f"Missing required environment variable: {name}")
    return val


def parse_modules(raw: str):
    items = [x.strip() for x in text(raw).split(",") if x.strip()]
    return items or ["Leads", "Contacts", "Deals"]

def build_parser() -> argparse.ArgumentParser:
    root = Path(__file__).resolve().parents[1]
    p = argparse.ArgumentParser(description="Import Zoho CRM into Dashboard Comercial Brasil DB.")
    p.add_argument("--workbook", default=str(root / "Dashboard Comercial Brasil.xlsx"))
    p.add_argument("--sheet", default="Comercial Brasil")
    p.add_argument("--header-row", type=int, default=2)
    p.add_argument("--data-start-row", type=int, default=3)
    p.add_argument("--leads-js", default=str(root / "src" / "dashboard" / "data" / "leads.js"))
    p.add_argument("--build-script", default=str(root / "tools" / "build-single-file.ps1"))
    p.add_argument("--modules", default=os.getenv("ZOHO_MODULES", "Leads,Contacts,Deals"))
    p.add_argument("--max-pages", type=int, default=20)
    p.add_argument("--per-page", type=int, default=200)
    p.add_argument("--overwrite", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--skip-build", action="store_true")
    p.add_argument("--no-backup", action="store_true")
    p.add_argument("--keep-backups", type=int, default=2)
    return p


def main() -> int:
    args = build_parser().parse_args()
    workbook = Path(args.workbook)
    leads_js = Path(args.leads_js)
    build_script = Path(args.build_script)

    if not workbook.exists():
        raise RuntimeError(f"Workbook not found: {workbook}")

    cid = need_env("ZOHO_CLIENT_ID")
    secret = need_env("ZOHO_CLIENT_SECRET")
    refresh = need_env("ZOHO_REFRESH_TOKEN")
    accounts_domain = os.getenv("ZOHO_ACCOUNTS_DOMAIN", "accounts.zoho.com").strip()
    api_domain = os.getenv("ZOHO_API_DOMAIN", "").strip()

    client = ZohoClient(cid, secret, refresh, accounts_domain, api_domain)
    client.authenticate()
    print(f"Connected to Zoho. API domain: {client.api_domain}")

    imported = []
    for module in parse_modules(args.modules):
        print(f"Fetching module: {module}")
        rows = client.get_records(module, args.max_pages, args.per_page)
        print(f"  - fetched {len(rows)} records")
        for r in rows:
            try:
                nr = normalize_record(module, r)
                if nr:
                    imported.append(nr)
            except Exception as e:
                print(f"  ! skipped malformed record in {module}: {e}")

    dedup = {}
    for r in imported:
        key = norm(r.get("zoho_key")) or "|".join([norm(r.get("Company")), norm(r.get("Contact_Company")), norm(r.get("Description"))])
        dedup[key] = r
    records = list(dedup.values())
    print(f"Normalized records: {len(records)}")

    wb = openpyxl.load_workbook(workbook)
    if args.sheet not in wb.sheetnames:
        raise RuntimeError(f"Sheet not found: {args.sheet}")
    ws = wb[args.sheet]

    cols = load_cols(ws, args.header_row)
    if "ID_Opportunity" not in cols:
        cols["ID_Opportunity"] = ensure_col(ws, args.header_row, "ID_Opportunity")

    for extra in ["Contact_Email", "Contact_Phone", "Contact_LinkedIn", "Contact_WhatsApp", "Zoho_Key", "Zoho_Module", "Zoho_Last_Sync"]:
        cols[extra] = ensure_col(ws, args.header_row, extra)
    cols = load_cols(ws, args.header_row)

    summary = sync_records(ws, cols, args.data_start_row, records, args.overwrite)
    print(f"Sync summary: created={summary['created']} updated={summary['updated']} skipped={summary['skipped']} total={summary['total']}")

    if args.dry_run:
        print("Dry-run mode: no files were saved.")
        return 0

    if not args.no_backup:
        backup = create_backup(workbook, args.keep_backups)
        print(f"Created backup: {backup.name}")

    wb.save(workbook)
    print(f"Workbook saved: {workbook}")

    n = export_leads_js(workbook, args.sheet, args.header_row, args.data_start_row, leads_js)
    print(f"Dashboard data exported ({n} rows): {leads_js}")

    if not args.skip_build and build_script.exists():
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(build_script)], check=True)
        print("Single-file dashboard rebuilt.")

    print("Zoho import completed successfully.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as e:
        print(f"ERROR: {e}")
        raise SystemExit(1)
