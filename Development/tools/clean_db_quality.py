#!/usr/bin/env python3
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def text(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    s = str(v).strip()
    return "" if s.lower() in {"nan", "none", "nat"} else s


def find_project_root() -> Path:
    here = Path(__file__).resolve().parent
    for p in [here, *here.parents]:
        if (p / "Apeiron_BR_Gestao_Comercial.xlsx").exists():
            return p
    return Path(__file__).resolve().parents[1]


def norm(v):
    return text(v).lower()


def to_int(v):
    s = text(v)
    if not s:
        return None
    try:
        return int(float(s))
    except Exception:
        return None


def backup_workbook(workbook_path: Path, backup_dir: Path, keep: int = 2) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    backup_path = backup_dir / f"Apeiron_BR_Gestao_Comercial_backup_{stamp}.xlsx"
    shutil.copyfile(workbook_path, backup_path)
    backups = sorted(
        backup_dir.glob("Apeiron_BR_Gestao_Comercial_backup_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    for old in backups[keep:]:
        try:
            old.unlink()
        except Exception:
            pass
    return backup_path


def mojibake_score(s: str) -> int:
    if not s:
        return 0
    score = 0
    score += s.count("Ã")
    score += s.count("Â")
    score += s.count("â")
    score += s.count("ð")
    score += s.count("�") * 3
    for tok in ["Ã©", "Ã£", "Ã¡", "Ã³", "Ãº", "Ã§", "â€”", "â€“", "â€œ", "â€"]:
        score += s.count(tok) * 2
    return score


def repair_mojibake(s: str) -> str:
    if not s:
        return s
    if not any(ch in s for ch in ("Ã", "Â", "â", "ð", "�")):
        return s

    candidates = [s]
    for enc in ("latin1", "cp1252"):
        try:
            candidates.append(s.encode(enc).decode("utf-8"))
        except Exception:
            pass

    best = min(candidates, key=mojibake_score)
    return best if mojibake_score(best) < mojibake_score(s) else s


def normalize_text_cells(wb) -> int:
    changed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                v = cell.value
                if not isinstance(v, str):
                    continue
                fixed = repair_mojibake(v)
                if fixed != v:
                    cell.value = fixed
                    changed += 1
    return changed


def header_map(ws, header_row):
    out = {}
    for c in range(1, ws.max_column + 1):
        h = text(ws.cell(header_row, c).value)
        if h:
            out[h] = c
    return out


def rows_by_id(ws, headers, start_row):
    out = {}
    id_col = headers.get("ID")
    if not id_col:
        return out
    for r in range(start_row, ws.max_row + 1):
        iid = to_int(ws.cell(r, id_col).value)
        if iid is not None:
            out[iid] = r
    return out


def completeness_score(ws, headers, row):
    score = 0
    for h, c in headers.items():
        if h == "ID":
            continue
        if text(ws.cell(row, c).value):
            score += 1
    return score


def merge_missing_values(ws, headers, target_row, source_row):
    merged = 0
    for h, c in headers.items():
        if h == "ID":
            continue
        tv = text(ws.cell(target_row, c).value)
        sv = text(ws.cell(source_row, c).value)
        if not tv and sv:
            ws.cell(target_row, c).value = sv
            merged += 1
    return merged


def delete_ids(ws, headers, start_row, ids):
    id_col = headers.get("ID")
    if not id_col:
        return 0
    rows = []
    wanted = set(ids)
    for r in range(start_row, ws.max_row + 1):
        iid = to_int(ws.cell(r, id_col).value)
        if iid in wanted:
            rows.append(r)
    for r in reversed(rows):
        ws.delete_rows(r, 1)
    return len(rows)


def export_leads_js(workbook_path: Path, out_js: Path):
    xls = pd.ExcelFile(workbook_path)
    ops = pd.read_excel(xls, sheet_name="Leads_Operations")
    det = pd.read_excel(xls, sheet_name="Leads_Details")

    det_map = {int(r["ID"]): r for _, r in det.dropna(subset=["ID"]).iterrows()}
    rows = []
    for _, r in ops.dropna(subset=["ID"]).iterrows():
        iid = int(r["ID"])
        d = det_map.get(iid, {})
        item = {
            "id": iid,
            "company": text(r.get("Company")),
            "name": text(r.get("Full Name")),
            "role": text(r.get("Role")),
            "status": text(r.get("Status")) or text(d.get("Status")),
            "priority": text(r.get("Priority")) or text(d.get("Priority")),
            "nextAction": text(r.get("Next Action")) or text(d.get("Next Action")),
            "nextActionDate": text(r.get("Next Action Date")) or text(d.get("Next Action Date")),
            "channel": text(r.get("Channel")),
            "owner": text(r.get("Lead Insight")) or text(d.get("Lead Insight")),
            "source": text(r.get("Source")) or text(d.get("Source")),
            "sharePointId": text(r.get("SharePoint_ID")) or text(d.get("SharePoint_ID")),
            "possibleDuplicate": text(r.get("Possible Duplicate")),
            "email": text(d.get("Email")),
            "phone": text(d.get("Phone")),
            "linkedin": text(d.get("LinkedIn")),
            "location": text(d.get("Location")),
            "department": text(d.get("Department")),
            "notes": text(d.get("Notes")),
            "accountSummary": text(d.get("Account_Summary")),
            "accountSector": text(d.get("Account_Sector")),
            "accountCityState": text(d.get("Account_City_State")),
            "accountSize": text(d.get("Account_Size")),
            "accountRevenue": text(d.get("Account_Revenue")),
            "accountWebsite": text(d.get("Account_Website")),
            "accountPhone": text(d.get("Account_Phone")),
        }
        rows.append(item)

    out_js.parent.mkdir(parents=True, exist_ok=True)
    out_js.write_text("window.INITIAL_DATA = " + json.dumps(rows, ensure_ascii=False) + ";\n", encoding="utf-8")
    return len(rows)


def main():
    root = find_project_root()
    support = root / "Support Files"
    workbook_path = root / "Apeiron_BR_Gestao_Comercial.xlsx"
    backup_dir = support / "_backup"
    out_dir = support / "_output"
    leads_js = root / "src" / "dashboard" / "data" / "leads.js"
    out_dir.mkdir(parents=True, exist_ok=True)

    backup_path = backup_workbook(workbook_path, backup_dir, keep=2)
    wb = load_workbook(workbook_path)

    ws_ops = wb["Leads_Operations"]
    ws_det = wb["Leads_Details"]
    contacts_name = next((s for s in wb.sheetnames if "Contacts" in s), None)
    if not contacts_name:
        raise RuntimeError("Contacts sheet not found.")
    ws_con = wb[contacts_name]

    ops_headers = header_map(ws_ops, 1)
    det_headers = header_map(ws_det, 1)
    con_headers = header_map(ws_con, 3)

    ops_rows = rows_by_id(ws_ops, ops_headers, 2)
    det_rows = rows_by_id(ws_det, det_headers, 2)
    con_rows = rows_by_id(ws_con, con_headers, 4)

    groups = {}
    for iid, row in ops_rows.items():
        k = f"{norm(ws_ops.cell(row, ops_headers.get('Full Name')).value)}|{norm(ws_ops.cell(row, ops_headers.get('Company')).value)}"
        if k == "|":
            continue
        groups.setdefault(k, []).append(iid)

    dup_groups = {k: ids for k, ids in groups.items() if len(ids) > 1}
    merged_cells = 0
    removed_ids = []

    for _, ids in dup_groups.items():
        candidates = []
        for iid in ids:
            score = 0
            if iid in ops_rows:
                score += completeness_score(ws_ops, ops_headers, ops_rows[iid])
            if iid in det_rows:
                score += completeness_score(ws_det, det_headers, det_rows[iid])
            if iid in con_rows:
                score += completeness_score(ws_con, con_headers, con_rows[iid])
            candidates.append((score, -iid, iid))
        candidates.sort(reverse=True)
        primary_id = candidates[0][2]
        to_remove = [iid for iid in ids if iid != primary_id]

        for iid in to_remove:
            if primary_id in ops_rows and iid in ops_rows:
                merged_cells += merge_missing_values(ws_ops, ops_headers, ops_rows[primary_id], ops_rows[iid])
            if primary_id in det_rows and iid in det_rows:
                merged_cells += merge_missing_values(ws_det, det_headers, det_rows[primary_id], det_rows[iid])
            if primary_id in con_rows and iid in con_rows:
                merged_cells += merge_missing_values(ws_con, con_headers, con_rows[primary_id], con_rows[iid])
        removed_ids.extend(to_remove)

    deleted_ops = delete_ids(ws_ops, ops_headers, 2, removed_ids)
    deleted_det = delete_ids(ws_det, det_headers, 2, removed_ids)
    deleted_con = delete_ids(ws_con, con_headers, 4, removed_ids)

    normalized_cells = normalize_text_cells(wb)

    wb.save(workbook_path)
    dashboard_rows = export_leads_js(workbook_path, leads_js)

    report_path = out_dir / "db_quality_cleanup_report.md"
    report_path.write_text(
        "\n".join(
            [
                "# DB Quality Cleanup Report",
                "",
                f"- generated_at: {datetime.now().isoformat(timespec='seconds')}",
                f"- workbook: `{workbook_path}`",
                f"- backup_created: `{backup_path}`",
                f"- duplicate_groups_found: {len(dup_groups)}",
                f"- duplicate_ids_removed: {len(removed_ids)}",
                f"- merged_cells_from_duplicates: {merged_cells}",
                f"- deleted_rows_ops: {deleted_ops}",
                f"- deleted_rows_details: {deleted_det}",
                f"- deleted_rows_contacts: {deleted_con}",
                f"- normalized_text_cells: {normalized_cells}",
                f"- dashboard_rows_after_export: {dashboard_rows}",
            ]
        ),
        encoding="utf-8",
    )

    print("backup:", backup_path)
    print("duplicate_groups:", len(dup_groups))
    print("removed_ids:", len(removed_ids))
    print("normalized_cells:", normalized_cells)
    print("dashboard_rows:", dashboard_rows)
    print("report:", report_path)


if __name__ == "__main__":
    main()
