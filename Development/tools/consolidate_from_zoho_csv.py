#!/usr/bin/env python3
import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def text(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    s = str(v).strip()
    return "" if s.lower() in {"nan", "none", "nat"} else s


def find_project_root() -> Path:
    here = Path(__file__).resolve().parent
    for p in [here, *here.parents]:
        if (p / "Apeiron_BR_Gestao_Comercial.xlsx").exists():
            return p
    return Path(__file__).resolve().parents[1]


def norm(v):
    return text(v).lower()


def clean_phone(v):
    return re.sub(r"[^0-9+]", "", text(v))


def channel_set_from_record(rec):
    out = set()
    if text(rec.get("Email")):
        out.add("Email")
    if clean_phone(rec.get("Phone")) or clean_phone(rec.get("Mobile")):
        out.add("Phone")
    for k, v in rec.items():
        lk = norm(k)
        if "linkedin" in lk and text(v):
            out.add("LinkedIn")
        if "whatsapp" in lk and text(v):
            out.add("WhatsApp")
    return out


def channel_label(ch_set):
    order = ["Email", "Phone", "LinkedIn", "WhatsApp"]
    vals = [x for x in order if x in ch_set]
    return ", ".join(vals) if vals else "No contact"


def first_last_name(rec):
    first = text(rec.get("First Name") or rec.get("First_Name"))
    last = text(rec.get("Last Name") or rec.get("Last_Name"))
    if not first and not last:
        full = text(rec.get("Full Name") or rec.get("Full_Name") or rec.get("Contact Name"))
        parts = [p for p in full.split() if p]
        if not parts:
            return "", "", ""
        if len(parts) == 1:
            return parts[0], "", parts[0]
        return parts[0], " ".join(parts[1:]), full
    full = " ".join([x for x in [first, last] if x]).strip()
    return first, last, full


def priority_from_record(rec):
    raw = norm(rec.get("Rating") or rec.get("Priority") or "")
    if "high" in raw or "hot" in raw:
        return "High"
    if "low" in raw or "cold" in raw:
        return "Low"
    return "Medium"


def status_from_record(rec):
    raw = norm(rec.get("Lead Status") or rec.get("Status") or "")
    if not raw:
        return "New"
    if "meeting" in raw or "schedule" in raw:
        return "Meeting Scheduled"
    if "proposal" in raw:
        return "Proposal Sent"
    if "contact" in raw:
        return "In Contact"
    if "client" in raw or "customer" in raw or "convert" in raw:
        return "Client"
    if "junk" in raw or "discard" in raw or "lost" in raw or "not interested" in raw:
        return "Discarded"
    return "New"


def backup_workbook(workbook_path, backup_dir, keep=2):
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    dst = os.path.join(backup_dir, f"Apeiron_BR_Gestao_Comercial_backup_{stamp}.xlsx")
    shutil.copyfile(workbook_path, dst)
    files = sorted(
        [f for f in os.listdir(backup_dir) if f.startswith("Apeiron_BR_Gestao_Comercial_backup_") and f.endswith(".xlsx")],
        reverse=True,
    )
    for old in files[keep:]:
        try:
            os.remove(os.path.join(backup_dir, old))
        except Exception:
            pass
    return dst


def header_map(ws, header_row):
    h = {}
    for c in range(1, ws.max_column + 1):
        v = text(ws.cell(header_row, c).value)
        if v:
            h[v] = c
    return h


def setv(ws, row, headers, col, value):
    c = headers.get(col)
    if c:
        ws.cell(row, c).value = value


def row_for_id(ws, headers, data_start, iid):
    id_col = headers.get("ID")
    if not id_col:
        return None
    for r in range(data_start, ws.max_row + 1):
        try:
            if int(float(ws.cell(r, id_col).value)) == int(iid):
                return r
        except Exception:
            continue
    return None


def next_id_from_sheet(ws, headers, data_start):
    id_col = headers.get("ID")
    if not id_col:
        return 1
    mx = 0
    for r in range(data_start, ws.max_row + 1):
        try:
            v = int(float(ws.cell(r, id_col).value))
            if v > mx:
                mx = v
        except Exception:
            pass
    return mx + 1


def export_leads_js(workbook_path, out_js):
    xls = pd.ExcelFile(workbook_path)
    ops = pd.read_excel(xls, sheet_name="Leads_Operations")
    det = pd.read_excel(xls, sheet_name="Leads_Details")
    det_map = {int(r["ID"]): r for _, r in det.dropna(subset=["ID"]).iterrows()}
    records = []
    for _, r in ops.dropna(subset=["ID"]).iterrows():
        iid = int(r["ID"])
        d = det_map.get(iid, {})
        item = {
            "id": iid,
            "company": text(r.get("Company")),
            "name": text(r.get("Full Name")),
            "role": text(r.get("Role")),
            "status": text(r.get("Status")) or text(d.get("Status")),
            "priority": text(r.get("Priority")) or text(d.get("Priority")),
            "nextAction": text(r.get("Next Action")) or text(d.get("Next Action")),
            "nextActionDate": text(r.get("Next Action Date")) or text(d.get("Next Action Date")),
            "channel": text(r.get("Channel")),
            "owner": text(r.get("Lead Insight")) or text(d.get("Lead Insight")),
            "source": text(r.get("Source")) or text(d.get("Source")),
            "sharePointId": text(r.get("SharePoint_ID")) or text(d.get("SharePoint_ID")),
            "possibleDuplicate": text(r.get("Possible Duplicate")),
            "email": text(d.get("Email")),
            "phone": text(d.get("Phone")),
            "linkedin": text(d.get("LinkedIn")),
            "location": text(d.get("Location")),
            "department": text(d.get("Department")),
            "notes": text(d.get("Notes")),
            "accountSummary": text(d.get("Account_Summary")),
            "accountSector": text(d.get("Account_Sector")),
            "accountCityState": text(d.get("Account_City_State")),
            "accountSize": text(d.get("Account_Size")),
            "accountRevenue": text(d.get("Account_Revenue")),
            "accountWebsite": text(d.get("Account_Website")),
            "accountPhone": text(d.get("Account_Phone")),
        }
        records.append(item)

    os.makedirs(os.path.dirname(out_js), exist_ok=True)
    with open(out_js, "w", encoding="utf-8") as f:
        f.write("window.INITIAL_DATA = " + json.dumps(records, ensure_ascii=False) + ";\n")
    return len(records)


def main():
    base = find_project_root()
    support = base / "Support Files"
    wb_path = str(base / "Apeiron_BR_Gestao_Comercial.xlsx")
    backup_dir = str(support / "_backup")
    zoho_dir = str((support / "Zoho data") if (support / "Zoho data").exists() else (base / "Zoho data"))
    out_dir = str(support / "_output")
    os.makedirs(out_dir, exist_ok=True)

    backup = backup_workbook(wb_path, backup_dir, keep=2)

    wb = load_workbook(wb_path)
    ws_ops = wb["Leads_Operations"]
    ws_det = wb["Leads_Details"]
    contacts_name = next((s for s in wb.sheetnames if "Contacts" in s), None)
    if not contacts_name:
        raise RuntimeError("Contacts sheet not found.")
    ws_con = wb[contacts_name]

    h_ops = header_map(ws_ops, 1)
    h_det = header_map(ws_det, 1)
    h_con = header_map(ws_con, 3)

    ops_df = pd.read_excel(wb_path, sheet_name="Leads_Operations")
    det_df = pd.read_excel(wb_path, sheet_name="Leads_Details")

    existing_emails = set(det_df.get("Email", pd.Series(dtype=str)).fillna("").astype(str).map(norm))
    existing_emails.discard("")
    existing_keys = set(
        ops_df.apply(lambda r: f"{norm(r.get('Full Name'))}|{norm(r.get('Company'))}", axis=1).tolist()
    )

    z_contacts = pd.read_csv(os.path.join(zoho_dir, "Contacts_2026_03_17.csv"), keep_default_na=False).to_dict("records")
    z_leads = pd.read_csv(os.path.join(zoho_dir, "Leads_2026_03_17.csv"), keep_default_na=False).to_dict("records")
    z_accounts = pd.read_csv(os.path.join(zoho_dir, "Accounts_2026_03_17.csv"), keep_default_na=False).to_dict("records")
    all_records = [("Contacts", r) for r in z_contacts] + [("Leads", r) for r in z_leads]
    account_by_name = {norm(a.get("Account Name")): a for a in z_accounts if norm(a.get("Account Name"))}

    next_id = max(
        next_id_from_sheet(ws_ops, h_ops, 2),
        next_id_from_sheet(ws_det, h_det, 2),
        next_id_from_sheet(ws_con, h_con, 4),
    )

    inserted = []
    skipped = 0
    for module, rec in all_records:
        company = text(rec.get("Account Name") if module == "Contacts" else rec.get("Company"))
        if not company:
            company = text(rec.get("Company") or rec.get("Account Name") or rec.get("Account Name.Account Name"))
        first, last, full = first_last_name(rec)
        if not full:
            continue

        email = norm(rec.get("Email"))
        key = f"{norm(full)}|{norm(company)}"
        if (email and email in existing_emails) or key in existing_keys:
            skipped += 1
            continue

        phone = clean_phone(rec.get("Phone"))
        mobile = clean_phone(rec.get("Mobile"))
        role = text(rec.get("Title") or rec.get("Designation"))
        department = text(rec.get("Department"))
        source = text(rec.get("Lead Source")) or text(rec.get("Source")) or f"Zoho CRM ({module})"
        status = status_from_record(rec)
        priority = priority_from_record(rec)
        ch_set = channel_set_from_record(rec)
        ch_label = channel_label(ch_set)
        location = ", ".join([x for x in [text(rec.get("City")), text(rec.get("State")), text(rec.get("Country"))] if x])
        account = account_by_name.get(norm(company), {})
        account_city = text(account.get("Billing City") or account.get("Shipping City"))
        account_state = text(account.get("Billing State") or account.get("Shipping State"))
        account_city_state = ", ".join([x for x in [account_city, account_state] if x])
        account_summary = text(account.get("Description"))
        account_sector = text(account.get("Industry"))
        account_size = text(account.get("Employees"))
        account_revenue = text(account.get("Annual Revenue"))
        account_website = text(account.get("Website"))
        account_phone = clean_phone(account.get("Phone"))

        # ops
        r_ops = ws_ops.max_row + 1
        setv(ws_ops, r_ops, h_ops, "ID", next_id)
        setv(ws_ops, r_ops, h_ops, "Company", company)
        setv(ws_ops, r_ops, h_ops, "Full Name", full)
        setv(ws_ops, r_ops, h_ops, "Role", role)
        setv(ws_ops, r_ops, h_ops, "Status", status)
        setv(ws_ops, r_ops, h_ops, "Priority", priority)
        setv(ws_ops, r_ops, h_ops, "Channel", ch_label)
        setv(ws_ops, r_ops, h_ops, "Channel Email", "Yes" if "Email" in ch_set else "")
        setv(ws_ops, r_ops, h_ops, "Channel Phone", "Yes" if "Phone" in ch_set else "")
        setv(ws_ops, r_ops, h_ops, "Channel LinkedIn", "Yes" if "LinkedIn" in ch_set else "")
        setv(ws_ops, r_ops, h_ops, "Channel WhatsApp", "Yes" if "WhatsApp" in ch_set else "")
        setv(ws_ops, r_ops, h_ops, "Source", source)
        setv(ws_ops, r_ops, h_ops, "Possible Duplicate", "")

        # details
        r_det = ws_det.max_row + 1
        setv(ws_det, r_det, h_det, "ID", next_id)
        setv(ws_det, r_det, h_det, "Company", company)
        setv(ws_det, r_det, h_det, "Full Name", full)
        setv(ws_det, r_det, h_det, "Role", role)
        setv(ws_det, r_det, h_det, "Department", department)
        setv(ws_det, r_det, h_det, "Email", email)
        setv(ws_det, r_det, h_det, "Phone", phone or mobile)
        setv(ws_det, r_det, h_det, "LinkedIn", text(rec.get("LinkedIn")))
        setv(ws_det, r_det, h_det, "Location", location)
        setv(ws_det, r_det, h_det, "Status", status)
        setv(ws_det, r_det, h_det, "Priority", priority)
        setv(ws_det, r_det, h_det, "Channel Email", "Yes" if "Email" in ch_set else "")
        setv(ws_det, r_det, h_det, "Channel Phone", "Yes" if "Phone" in ch_set else "")
        setv(ws_det, r_det, h_det, "Channel LinkedIn", "Yes" if "LinkedIn" in ch_set else "")
        setv(ws_det, r_det, h_det, "Channel WhatsApp", "Yes" if "WhatsApp" in ch_set else "")
        setv(ws_det, r_det, h_det, "Source", source)
        setv(ws_det, r_det, h_det, "Created Date", datetime.now().strftime("%Y-%m-%d"))
        setv(ws_det, r_det, h_det, "Account_Summary", account_summary)
        setv(ws_det, r_det, h_det, "Account_Sector", account_sector)
        setv(ws_det, r_det, h_det, "Account_City_State", account_city_state)
        setv(ws_det, r_det, h_det, "Account_Size", account_size)
        setv(ws_det, r_det, h_det, "Account_Revenue", account_revenue)
        setv(ws_det, r_det, h_det, "Account_Website", account_website)
        setv(ws_det, r_det, h_det, "Account_Phone", account_phone)

        # contacts
        r_con = ws_con.max_row + 1
        setv(ws_con, r_con, h_con, "ID", next_id)
        setv(ws_con, r_con, h_con, "Company", company)
        setv(ws_con, r_con, h_con, "First Name", first)
        setv(ws_con, r_con, h_con, "Last Name", last)
        setv(ws_con, r_con, h_con, "Role", role)
        setv(ws_con, r_con, h_con, "Department", department)
        setv(ws_con, r_con, h_con, "Email", email)
        setv(ws_con, r_con, h_con, "Phone", phone or mobile)
        setv(ws_con, r_con, h_con, "LinkedIn", text(rec.get("LinkedIn")))
        setv(ws_con, r_con, h_con, "Location", location)
        setv(ws_con, r_con, h_con, "Status", status)
        setv(ws_con, r_con, h_con, "Priority", priority)
        setv(ws_con, r_con, h_con, "Channel Email", "Yes" if "Email" in ch_set else "")
        setv(ws_con, r_con, h_con, "Channel Phone", "Yes" if "Phone" in ch_set else "")
        setv(ws_con, r_con, h_con, "Channel LinkedIn", "Yes" if "LinkedIn" in ch_set else "")
        setv(ws_con, r_con, h_con, "Channel WhatsApp", "Yes" if "WhatsApp" in ch_set else "")
        setv(ws_con, r_con, h_con, "Source", source)
        setv(ws_con, r_con, h_con, "Created Date", datetime.now().strftime("%Y-%m-%d"))
        setv(ws_con, r_con, h_con, "Account_Summary", account_summary)
        setv(ws_con, r_con, h_con, "Account_Sector", account_sector)
        setv(ws_con, r_con, h_con, "Account_City_State", account_city_state)
        setv(ws_con, r_con, h_con, "Account_Size", account_size)
        setv(ws_con, r_con, h_con, "Account_Revenue", account_revenue)
        setv(ws_con, r_con, h_con, "Account_Website", account_website)
        setv(ws_con, r_con, h_con, "Account_Phone", account_phone)

        existing_keys.add(key)
        if email:
            existing_emails.add(email)
        inserted.append(
            {
                "ID": next_id,
                "module": module,
                "company": company,
                "name": full,
                "email": email,
                "phone": phone or mobile,
                "source": source,
            }
        )
        next_id += 1

    wb.save(wb_path)

    inserted_csv = os.path.join(out_dir, "zoho_inserted_rows.csv")
    pd.DataFrame(inserted).to_csv(inserted_csv, index=False)

    leads_js = os.path.join(base, "src", "dashboard", "data", "leads.js")
    total_dashboard = export_leads_js(wb_path, leads_js)

    report_path = os.path.join(out_dir, "consolidation_report.md")
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("# Consolidation Report - Zoho CSV -> DB\n\n")
        f.write(f"- generated_at: {datetime.now().isoformat(timespec='seconds')}\n")
        f.write(f"- workbook: `{wb_path}`\n")
        f.write(f"- backup_created: `{backup}`\n")
        f.write(f"- inserted_rows: {len(inserted)}\n")
        f.write(f"- skipped_existing: {skipped}\n")
        f.write(f"- inserted_export: `{inserted_csv}`\n")
        f.write(f"- dashboard_rows_after_export: {total_dashboard}\n")

    print("backup:", backup)
    print("inserted:", len(inserted))
    print("skipped:", skipped)
    print("inserted_csv:", inserted_csv)
    print("report:", report_path)
    print("dashboard_rows:", total_dashboard)


if __name__ == "__main__":
    main()
