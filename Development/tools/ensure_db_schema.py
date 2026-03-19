#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def find_project_root() -> Path:
    here = Path(__file__).resolve().parent
    for p in [here, *here.parents]:
        if (p / "Apeiron_BR_Gestao_Comercial.xlsx").exists():
            return p
    return Path(__file__).resolve().parents[1]


def headers_row_1(ws):
    out = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        out.append("" if v is None else str(v).strip())
    return out


def ensure_column(ws, header_name: str) -> bool:
    headers = headers_row_1(ws)
    if header_name in headers:
        return False
    col = ws.max_column + 1
    ws.cell(1, col).value = header_name
    return True


def main() -> None:
    parser = argparse.ArgumentParser(description="Ensure DB schema columns for enrichment.")
    parser.add_argument("--workbook", type=Path, default=None)
    args = parser.parse_args()

    root = find_project_root()
    workbook = args.workbook or (root / "Apeiron_BR_Gestao_Comercial.xlsx")
    wb = load_workbook(workbook)

    changes = []
    if "Leads_Details" in wb.sheetnames:
        ws = wb["Leads_Details"]
        if ensure_column(ws, "Account_CNPJ"):
            changes.append("Leads_Details.Account_CNPJ")
        if ensure_column(ws, "Account_Brand"):
            changes.append("Leads_Details.Account_Brand")
    if "Accounts" in wb.sheetnames:
        ws = wb["Accounts"]
        if ensure_column(ws, "CNPJ"):
            changes.append("Accounts.CNPJ")
        if ensure_column(ws, "Brand"):
            changes.append("Accounts.Brand")

    if changes:
        wb.save(workbook)

    report = {
        "workbook": str(workbook),
        "changes": changes,
        "changed": bool(changes),
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
